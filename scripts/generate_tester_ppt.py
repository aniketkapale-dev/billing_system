"""
Generate QA/Tester PowerPoint for DailyBasket Billing System.
Run: python scripts/generate_tester_ppt.py
Output: docs/DailyBasket_Billing_System_Tester_Guide.pptx
"""

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

OUTPUT = Path(__file__).resolve().parent.parent / "docs" / "DailyBasket_Billing_System_Tester_Guide.pptx"

NAVY = RGBColor(0x1A, 0x23, 0x7E)
BLUE = RGBColor(0x15, 0x65, 0xC0)
ORANGE = RGBColor(0xE6, 0x51, 0x00)
DARK = RGBColor(0x21, 0x21, 0x21)
GRAY = RGBColor(0x61, 0x61, 0x61)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xE8, 0xEA, 0xF6)
RED = RGBColor(0xC6, 0x28, 0x28)


def set_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def title_slide(prs, title, subtitle=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, NAVY)
    b = slide.shapes.add_textbox(Inches(0.6), Inches(2.0), Inches(8.8), Inches(1.4))
    p = b.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER
    if subtitle:
        b2 = slide.shapes.add_textbox(Inches(0.6), Inches(3.6), Inches(8.8), Inches(1.6))
        p2 = b2.text_frame.paragraphs[0]
        p2.text = subtitle
        p2.font.size = Pt(18)
        p2.font.color.rgb = WHITE
        p2.alignment = PP_ALIGN.CENTER


def section_slide(prs, text):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, BLUE)
    b = slide.shapes.add_textbox(Inches(0.6), Inches(3.0), Inches(8.8), Inches(1.0))
    p = b.text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER


def content_slide(prs, title, bullets, note="", title_color=BLUE, font_size=15):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, LIGHT)
    bar = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(10), Inches(0.95))
    bar.fill.solid()
    bar.fill.fore_color.rgb = title_color
    bar.line.fill.background()
    tb = slide.shapes.add_textbox(Inches(0.45), Inches(0.12), Inches(9.1), Inches(0.75))
    tp = tb.text_frame.paragraphs[0]
    tp.text = title
    tp.font.size = Pt(22)
    tp.font.bold = True
    tp.font.color.rgb = WHITE
    body = slide.shapes.add_textbox(Inches(0.45), Inches(1.15), Inches(9.1), Inches(5.8))
    tf = body.text_frame
    tf.word_wrap = True
    for i, item in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if isinstance(item, tuple):
            text, level = item
            p.text = text
            p.level = level
        else:
            p.text = item
            p.level = 0
        p.font.size = Pt(font_size - (2 * p.level))
        p.font.color.rgb = DARK
        p.space_after = Pt(6)
    if note:
        nb = slide.shapes.add_textbox(Inches(0.45), Inches(6.85), Inches(9.1), Inches(0.5))
        np = nb.text_frame.paragraphs[0]
        np.text = note
        np.font.size = Pt(12)
        np.font.italic = True
        np.font.color.rgb = RED if "FAIL" in note or "Block" in note else GRAY


def table_slide(prs, title, headers, rows):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, LIGHT)
    bar = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(10), Inches(0.95))
    bar.fill.solid()
    bar.fill.fore_color.rgb = BLUE
    bar.line.fill.background()
    tb = slide.shapes.add_textbox(Inches(0.45), Inches(0.12), Inches(9.1), Inches(0.75))
    tb.text_frame.paragraphs[0].text = title
    tb.text_frame.paragraphs[0].font.size = Pt(20)
    tb.text_frame.paragraphs[0].font.bold = True
    tb.text_frame.paragraphs[0].font.color.rgb = WHITE

    cols = len(headers)
    nrows = len(rows) + 1
    left, top, width, height = Inches(0.35), Inches(1.1), Inches(9.3), Inches(0.35 * nrows)
    table = slide.shapes.add_table(nrows, cols, left, top, width, height).table
    for c, h in enumerate(headers):
        cell = table.cell(0, c)
        cell.text = h
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(11)
            p.font.bold = True
            p.font.color.rgb = WHITE
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
    for r, row in enumerate(rows, 1):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.text = str(val)
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(10)
                p.font.color.rgb = DARK


def test_case_slide(prs, module, cases):
    content_slide(
        prs,
        f"Test Cases — {module}",
        [f"TC-{i+1}: {tc}" for i, tc in enumerate(cases)],
        font_size=14,
    )


def build():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    # ===== INTRO =====
    title_slide(
        prs,
        "DailyBasket Billing System",
        "QA / Tester Guide\nDeveloper Perspective — Validations, Flows, Restrictions & Test Scenarios",
    )

    content_slide(
        prs,
        "Document Purpose",
        [
            "Complete testing reference for QA engineers and testers",
            "Covers every module: API endpoints, UI flows, validations, and business rules",
            "Includes expected error messages, restrictions, edge cases, and test case IDs",
            "Stack: Django REST API + JWT auth + JavaScript frontend",
            "All business APIs require: JWT token + X-Business-Id header + tab permission",
        ],
    )

    content_slide(
        prs,
        "Test Environment Setup",
        [
            "Base URL: http://127.0.0.1:8000 (local dev)",
            "Login page: /login/  |  Register: /register/",
            "Dashboard: /dashboard/ (after login as Business Owner/Staff)",
            "API prefix: /api/",
            "Required headers for business APIs:",
            ("Authorization: Bearer <access_token>", 1),
            ("X-Business-Id: <active_business_id>", 1),
            "Roles: Super Admin (platform) | Business Owner | Business Staff (tab-limited)",
            "Staff access controlled by role → allowed_tabs (menu permissions)",
        ],
    )

    section_slide(prs, "1. Authentication & Access Control")

    table_slide(
        prs,
        "Auth API Endpoints — /api/auth/",
        ["Method", "Route", "Auth", "Purpose"],
        [
            ("POST", "/register/", "Public", "Business owner registration"),
            ("POST", "/login/", "Public", "Login via email OR 10-digit mobile"),
            ("POST", "/logout/", "JWT", "Client-side token discard"),
            ("POST", "/refresh/", "Public", "Refresh access token"),
            ("POST", "/forgot-password/", "Public", "Stub — no email sent"),
            ("POST", "/reset-password/", "Public", "Reset via JWT reset token (15 min)"),
            ("POST", "/change-password/", "JWT", "Change password"),
            ("GET/PATCH", "/me/", "JWT", "Profile read/update"),
        ],
    )

    content_slide(
        prs,
        "Auth — Field Validations",
        [
            "Register: full_name (required, max 150), mobile (required, 10 digits, unique)",
            "Register: email (optional, valid format, unique), password (min 8 chars)",
            "Login: field named 'email' accepts email OR 10-digit mobile + password",
            "Change password: current_password + new_password (min 8)",
            "Profile update: full_name non-empty if sent; mobile/email unique",
        ],
    )

    content_slide(
        prs,
        "Auth — Business Rules & Restrictions",
        [
            "New registration → is_active=false until Super Admin approval",
            "Login with inactive account → BLOCKED with approval-wait message",
            "Login ID: 10-digit numeric → mobile lookup; else → email (case-insensitive)",
            "JWT refresh requires non-empty refresh token",
            "Frontend: login clears tokens on page load/back-navigation",
            "Post-login routing: Super Admin → /superadmin/ | Business user → /dashboard/",
        ],
        note="Expected errors: 'Invalid email/mobile or password.' | 'Password must be at least 8 characters.'",
    )

    test_case_slide(
        prs,
        "Authentication",
        [
            "Register with valid data → account inactive → login fails until approved",
            "Login with 10-digit mobile vs email — both must work",
            "Password with 7 chars → FAIL 'Password must be at least 8 characters.'",
            "Duplicate mobile on register → FAIL",
            "Wrong current password on change → FAIL 'Current password is incorrect.'",
            "Access /api/products/ without JWT → 401",
            "Access with JWT but missing X-Business-Id → business error",
        ],
    )

    section_slide(prs, "2. Business Profile & Settings")

    table_slide(
        prs,
        "Business API — /api/businesses/",
        ["Method", "Route", "Notes"],
        [
            ("GET", "/", "List owner's businesses"),
            ("POST", "/", "Create business"),
            ("GET/PUT/PATCH/DELETE", "/<id>/", "CRUD"),
        ],
    )

    content_slide(
        prs,
        "Business Profile — Validations",
        [
            "business_name: REQUIRED, unique per owner (active records)",
            "phone: optional → if set, exactly 10 digits",
            "email: optional, Django email validator",
            "gst_number, address: optional, trimmed",
            "logo: file upload; clear_logo=true removes logo on update",
            "Frontend error: 'Business name is required.' | 'No active business selected.'",
        ],
    )

    table_slide(
        prs,
        "Settings APIs",
        ["Module", "Base Route", "Tab Permission"],
        [
            ("Tax", "/api/settings/taxes/", "settings-tax"),
            ("Invoice", "/api/settings/invoice-settings/", "settings-invoice"),
            ("Barcode", "/api/settings/barcodes/", "settings-barcode (+ stock-in, products)"),
        ],
    )

    content_slide(
        prs,
        "Tax Settings — Validations & Rules",
        [
            "key: REQUIRED, max 100, unique per business",
            "value: REQUIRED, 0–100 (decimal)",
            "is_active: boolean",
            "Errors: 'A tax with this key already exists.' | 'Tax value must be between 0 and 100.'",
            "Soft delete + POST /<id>/restore/ supported",
        ],
    )

    content_slide(
        prs,
        "Invoice Settings — Validations & Rules",
        [
            "year: REQUIRED, 2000–2100",
            "prefix, suffix: optional max 50; combo unique per business+year",
            "counter: integer ≥ 0; sets current_counter on create",
            "end_counter: date — series valid while end_counter >= today",
            "terms_conditions: optional text; qr_image upload; clear_qr=true removes QR",
            "Invoice number format: {prefix}/{counter}/{suffix} (empty parts omitted)",
            "On sale finalize: current_counter incremented atomically (select_for_update)",
            "Series rollover: if expired → auto year+1 series (max year 2100)",
        ],
        note="Error: 'Invoice settings with this year, prefix, and suffix already exist.'",
    )

    test_case_slide(
        prs,
        "Settings",
        [
            "Create tax with value 101 → FAIL",
            "Duplicate tax key same business → FAIL",
            "Create invoice setting year 2026, prefix INV → verify number format on sale",
            "Finalize sale → counter increments by 1",
            "Set end_counter in past → verify series rollover on next sale",
            "Upload QR image on invoice setting → appears on print preview",
        ],
    )

    section_slide(prs, "3. Catalog Master Data")

    table_slide(
        prs,
        "Catalog APIs — /api/catalog/",
        ["Resource", "Route", "Tab", "Unique (per business)"],
        [
            ("Units", "/units/", "products-units", "name, short_name"),
            ("Categories", "/categories/", "products-categories", "name"),
            ("Brands", "/brands/", "products-brands", "name"),
            ("Manufacturers", "/manufacturers/", "products", "name"),
            ("Vendors", "/vendors/", "stock-in", "name"),
            ("Payment Types", "/payment-types/", "purchases", "name"),
        ],
    )

    content_slide(
        prs,
        "Catalog — Validations (All Entities)",
        [
            "All: name REQUIRED; max 150 chars (Unit name 100, short_name 20)",
            "Unit: name + short_name both required and unique",
            "Vendor: name required — used in Stock In",
            "Payment type: name required — used in Sales payments",
            "All support: GET/POST list, GET/PUT/PATCH/DELETE detail, POST /<id>/restore/",
            "Soft delete: is_deleted=true, is_active=false; default lists exclude deleted",
            "include_deleted=true query param shows soft-deleted rows",
        ],
    )

    test_case_slide(
        prs,
        "Catalog",
        [
            "Create duplicate category name → FAIL",
            "Create unit with duplicate short_name → FAIL",
            "Soft delete vendor → restore → verify reappears in Stock In dropdown",
            "Staff without products-categories tab → 403 on category API",
            "Create vendor from Stock In quick-add modal → name required validation",
        ],
    )

    section_slide(prs, "4. Products")

    table_slide(
        prs,
        "Products API — /api/products/  (Tab: products)",
        ["Method", "Route", "Notes"],
        [
            ("GET", "/", "Filters: category_id, unit_id, search, ordering"),
            ("POST", "/", "Create product"),
            ("GET", "/next-sku/", "Auto SKU: SKU-000001 pattern"),
            ("GET/PUT/PATCH/DELETE", "/<id>/", "CRUD — soft delete"),
        ],
    )

    content_slide(
        prs,
        "Products — Field Validations",
        [
            "name: required, max 150",
            "sku: required on create, max 50, UNIQUE per business (non-deleted)",
            "category_id: REQUIRED on create — active, same business",
            "unit_id: REQUIRED on create — FK PROTECT (cannot delete unit in use)",
            "actual_price, mrp, sale_price: min 0",
            "purchase_price: READ-ONLY — computed as actual_price × (1 + tax_rate/100)",
            "quantity (opening stock): min 0; if > 0 → actual_price must be > 0",
            "brand_id, manufacturer_id, tax_id, tax_ids: optional",
        ],
    )

    content_slide(
        prs,
        "Products — Business Rules & Restrictions",
        [
            "Opening stock > 0 → creates OPEN batch + inventory stock record",
            "SKU auto-gen: SKU-###### if empty on update",
            "PRODUCT WITH SALES LOCK: if product appears on any sale (PurchaseItem qty > 0):",
            ("→ Edit allowed ONLY for mrp and sale_price fields", 1),
            ("→ Delete BLOCKED — UI hides edit/delete; API returns error", 1),
            "Quantity increase → new ADJ opening batch + add stock",
            "Quantity decrease → FIFO consume batches; fail if insufficient stock",
            "Product MRP synced from max batch MRP after stock-in/opening batches",
        ],
        note="Error: 'This product has sales records and cannot be edited or deleted.'",
    )

    test_case_slide(
        prs,
        "Products",
        [
            "Create product without category → FAIL 'Category is required.'",
            "Duplicate SKU same business → FAIL",
            "Create with opening stock 10 but price 0 → FAIL",
            "GET /next-sku/ → returns next available SKU number",
            "Sell product once → try edit name → FAIL; edit MRP → PASS",
            "Reduce product qty below available batch stock → FAIL with current stock in message",
            "Delete unit referenced by product → FAIL (PROTECT)",
        ],
    )

    section_slide(prs, "5. Customers")

    table_slide(
        prs,
        "Customers API — /api/customers/  (Tab: customers)",
        ["Method", "Route"],
        [
            ("GET/POST", "/"),
            ("GET/PUT/PATCH/DELETE", "/<id>/"),
            ("POST", "/<id>/restore/"),
        ],
    )

    content_slide(
        prs,
        "Customers — Validations & Rules",
        [
            "name: REQUIRED, unique per business (case-insensitive)",
            "mobile: REQUIRED, exactly 10 digits",
            "address: REQUIRED",
            "pin_code: REQUIRED, exactly 6 digits",
            "email: optional, lowercased",
            "company_name, company_mobile, gst_number, business_address: optional",
            "company_mobile: if set → 10-digit validation",
            "Frontend: company name required when 'Add company' checkbox selected",
            "Search fields: name, company_name, mobile, email, pin_code",
        ],
        note="Errors: 'Enter a valid 6-digit pin code.' | 'A customer with this name already exists.'",
    )

    test_case_slide(
        prs,
        "Customers",
        [
            "Create without pin_code → FAIL",
            "Pin code 12345 (5 digits) → FAIL",
            "Duplicate customer name (case different) → FAIL",
            "Create with company → verify Company column in list + invoice print shows Mob. No.",
            "Mobile with letters → stripped/rejected",
        ],
    )

    section_slide(prs, "6. Stock In (Purchase Invoices)")

    table_slide(
        prs,
        "Stock In API — /api/invoicing/purchase-invoices/  (Tab: stock-in)",
        ["Method", "Route", "Notes"],
        [
            ("GET", "/", "Filters: date_from, date_to, search, invoice_number"),
            ("POST", "/", "Create with items + optional attachment (multipart)"),
            ("GET", "/<id>/", "Detail with batch_lines"),
            ("PATCH", "/<id>/", "HEADER ONLY — items NOT editable"),
            ("DELETE", "/<id>/", "Soft delete + stock reversal"),
        ],
    )

    content_slide(
        prs,
        "Stock In — Line Item Validations",
        [
            "invoice_number: REQUIRED, max 50, UNIQUE per business",
            "items: minimum 1 item required",
            "product_id: must exist in business",
            "quantity: must be > 0",
            "purchase_price: must be ≥ 0",
            "discount, tax: optional, default 0",
            "batch_number: optional max 50; auto-gen format B{YYYYMMDD}-{vendorId|0}-{priceInPaise}",
            "vendor_id: optional; must exist if provided",
            "mrp: optional min 0; auto-bump if purchase_price > max MRP",
            "attachment: PDF/JPG/JPEG/PNG/WEBP/GIF, max 5 MB",
        ],
    )

    content_slide(
        prs,
        "Stock In — Business Rules & Restrictions",
        [
            "Each line → creates InventoryBatch + increases aggregate stock",
            "Line total: (qty × purchase_price) - discount",
            "Grand total: subtotal - invoice_discount + invoice_tax",
            "MRP logic: if purchase_price > current max batch MRP → batch MRP = purchase_price",
            "DELETE restriction: BLOCKED if any batch stock already sold (available < purchased)",
            "On delete: reverses remaining batch qty from aggregate stock",
            "PATCH on edit: only header (number, date, remarks, attachment) — NOT line items",
            "Barcode: auto-fills if exists; prompts generate if missing; field is readonly",
        ],
        note="Error: 'Cannot delete this purchase because some stock has already been sold or used.'",
    )

    test_case_slide(
        prs,
        "Stock In",
        [
            "Create with duplicate invoice_number → FAIL",
            "Create with 0 items → FAIL 'At least one invoice item is required.'",
            "Upload 6MB attachment → FAIL 'File size must be 5 MB or less.'",
            "Upload .exe file → FAIL 'Only PDF and image files are allowed.'",
            "Save stock-in → verify inventory qty increased + batch created",
            "Partially sell batch stock → try delete stock-in → FAIL",
            "Edit stock-in → verify line items cannot be changed (PATCH header only)",
            "Change vendor/cost → verify batch number auto-regenerates (unless manually edited)",
        ],
    )

    section_slide(prs, "7. Sales / Create Invoice")

    content_slide(
        prs,
        "Important: Sales API Mapping",
        [
            "UI label 'Create Invoice' / 'Sales' maps to API: /api/purchases/",
            "Model name: Purchase (historical naming — this is CUSTOMER SALES, not supplier purchase)",
            "Tab permission: purchases",
            "DELETE /<id>/ → 405 NOT SUPPORTED — use mark-cancelled instead",
        ],
        title_color=ORANGE,
    )

    table_slide(
        prs,
        "Sales API — /api/purchases/",
        ["Method", "Route", "Purpose"],
        [
            ("GET", "/", "List; filter: date_from, date_to, invoice_status"),
            ("POST", "/", "Create sale (final or draft)"),
            ("GET/PATCH", "/<id>/", "Detail / update"),
            ("POST", "/<id>/finalize/", "Finalize draft → deduct stock"),
            ("POST", "/<id>/mark-paid/", "Mark fully paid"),
            ("GET/POST", "/<id>/payments/", "List / record payment"),
            ("POST", "/<id>/mark-cancelled/", "Cancel finalized or delete draft"),
        ],
    )

    content_slide(
        prs,
        "Sales — Header Validations",
        [
            "customer_id: REQUIRED",
            "invoice_setting_id: REQUIRED unless is_draft=true",
            "items: REQUIRED, minimum 1 line",
            "is_draft: default false",
            "due_date: REQUIRED for unpaid non-draft sales",
            "payment_amount: optional partial payment on create/finalize",
            "payment_type_id, reference_no: optional",
        ],
    )

    content_slide(
        prs,
        "Sales — Line Item Validations",
        [
            "product_id: required",
            "quantity: must be > 0",
            "unit_price: must be ≥ 0",
            "discount_type: 'percent' or 'amount'",
            "sale_tax_ids: optional list of tax IDs",
            "Frontend: MRP required per row; qty cannot exceed available stock",
            "Frontend label: 'Salon Price' (selling price field)",
            "Same product in multiple rows → quantities aggregated for stock check",
        ],
    )

    table_slide(
        prs,
        "Draft vs Finalized Sale — Critical Rules",
        ["State", "Stock", "Invoice #", "Payments", "Edit Items"],
        [
            ("Draft", "NOT deducted", "Empty", "Not allowed", "Full edit allowed"),
            ("Finalized", "FIFO deduct + aggregate deduct", "Allocated from counter", "Allowed", "Items IMMUTABLE"),
            ("Cancelled", "Restored via batch consumption", "Kept for audit", "Blocked", "N/A"),
        ],
    )

    content_slide(
        prs,
        "Sales — Payment Rules",
        [
            "Draft must be finalized before recording payments",
            "Payment amount must be > 0 and cannot exceed pending bill",
            "Payment date cannot be before last payment date",
            "Partial payment requires next_due_date ≥ payment date",
            "Full payment clears due date",
            "Cancelled invoices cannot receive payments",
            "payment_status computed: draft / cancelled / unpaid / partial / paid",
        ],
    )

    content_slide(
        prs,
        "Sales — Cancel & Edit Restrictions",
        [
            "Cancel finalized: cancellation_reason REQUIRED",
            "cancellation_date must be ≥ invoice_date",
            "Cancel restores FIFO batch consumptions + aggregate stock",
            "Cancel draft: soft-deletes purchase; may release last allocated invoice number",
            "Edit finalized: line items in PATCH body → 400 error",
            "Paid invoices: due_date changes ignored on edit",
            "Print preview: Terms & Conditions editable per print (localStorage override)",
        ],
        note="Errors: 'Sale line items cannot be changed after the invoice is created.'",
    )

    test_case_slide(
        prs,
        "Sales / Create Invoice",
        [
            "Save draft → verify stock NOT reduced; no invoice number assigned",
            "Finalize draft with insufficient stock → FAIL at finalize step",
            "Create final sale → stock reduced FIFO (oldest batch first)",
            "Multi-row same product exceeding total stock → FAIL",
            "Partial payment without next_due_date → FAIL",
            "Payment amount > pending bill → FAIL",
            "Cancel finalized with empty reason → FAIL",
            "Cancel with date before invoice date → FAIL",
            "PATCH finalized sale with items array → 400",
            "Verify invoice print: Mob. No. label, company contact, GST breakdown",
        ],
    )

    section_slide(prs, "8. Inventory & FIFO Batches")

    table_slide(
        prs,
        "Inventory APIs",
        ["Module", "Route", "Tab"],
        [
            ("Stock list", "/api/inventory/", "inventory"),
            ("Profit summary", "/api/inventory/profit-summary/", "inventory"),
            ("Batches", "/api/invoicing/batches/", "inventory"),
        ],
    )

    content_slide(
        prs,
        "FIFO & Batch Rules",
        [
            "Batch created on: Stock In line, Product opening stock, Product qty increase",
            "Batch fields: purchased_quantity, available_quantity, purchase_price, mrp, batch_number",
            "FIFO consume order: created_at ASC, then id ASC (oldest first)",
            "On sale finalize: BatchConsumption records linked to PurchaseItem",
            "Profit per slice: (selling_price - batch.purchase_price) × qty",
            "If aggregate stock > 0 but no batches → auto-creates OPEN batch on sale check",
            "Inventory list default filter: in_stock=true (qty > 0)",
        ],
    )

    test_case_slide(
        prs,
        "Inventory & FIFO",
        [
            "Stock-in 2 batches same product at different costs → sell → verify oldest consumed first",
            "Check batch_lines on purchase invoice detail API",
            "Verify profit-summary KPIs after sales",
            "Filter inventory in_stock=false → shows zero-stock products",
            "Cancel sale → verify batch available_quantity restored",
        ],
    )

    section_slide(prs, "9. Barcode Management")

    table_slide(
        prs,
        "Barcode API — /api/settings/barcodes/",
        ["Method", "Route", "Notes"],
        [
            ("GET/POST", "/", "Paginated: default 100/page, max 5000"),
            ("POST", "/bulk-generate/", "One barcode per SKU"),
            ("GET/PUT/PATCH/DELETE", "/<id>/", "CRUD + soft delete"),
            ("POST", "/<id>/restore/", "Restore deleted"),
        ],
    )

    content_slide(
        prs,
        "Barcode — Validations & Restrictions",
        [
            "ONE barcode per SKU — enforced backend + frontend",
            "value: REQUIRED, max 100, unique per business",
            "Bulk generate: product must have non-empty SKU",
            "Value format: YYYYMMDD + 4-digit random suffix",
            "DELETE BLOCKED if product used in any purchase invoice (stock-in)",
            "Cannot reassign barcode to different product if already assigned",
            "Generate from Stock In: opens modal if product has no barcode",
            "Reprint tab: SKU dropdown (not barcode number input); qty 1–500",
            "Barcode field on Stock In: readonly text input (no manual + button)",
        ],
        note="Error: 'This barcode cannot be deleted because the product was used in a purchase invoice.'",
    )

    test_case_slide(
        prs,
        "Barcodes",
        [
            "Bulk generate for SKU → success; generate again same SKU → FAIL",
            "Generate barcode without SKU on product → FAIL",
            "Delete barcode after stock-in uses product → FAIL; UI hides delete button",
            "Reprint with invalid qty 0 or 501 → validation error",
            "Stock-in product without barcode → modal auto-opens for generation",
            "Verify barcode auto-fills on stock-in line after generation",
        ],
    )

    section_slide(prs, "10. Global API Conventions")

    content_slide(
        prs,
        "Pagination & List Response",
        [
            "Query params: page (default 1), page_size (default 10, max 100; barcodes max 5000)",
            "search: case-insensitive across module search_fields",
            "ordering: field name; prefix '-' for descending",
            "include_deleted=true: shows soft-deleted records",
            "Response: { isSuccess, message, data: { items: [], pagination: { count, page, page_size, total_pages, has_next, has_previous } } }",
        ],
    )

    content_slide(
        prs,
        "Soft Delete Pattern",
        [
            "DELETE /<id>/ → sets is_deleted=true, is_active=false, deleted_at, deleted_by",
            "Default queries exclude soft-deleted (ActiveManager)",
            "POST /<id>/restore/ on: catalog, customers, settings entities",
            "Active duplicate name check still applies after restore",
        ],
    )

    content_slide(
        prs,
        "Cross-Module Restrictions Summary",
        [
            "All APIs scoped by X-Business-Id header",
            "Tab permission required per viewset (role → allowed_tabs)",
            "Product with sales → edit MRP/sale_price only; no delete",
            "FIFO on: sale finalize, product qty decrease, cancel restores stock",
            "Invoice numbering: shared atomic counter on InvoiceSetting",
            "Unit FK PROTECT → cannot delete unit used by products",
            "Sale DELETE → 405; use mark-cancelled",
            "Stock-in DELETE → only if no batch stock consumed",
        ],
        title_color=ORANGE,
    )

    section_slide(prs, "11. End-to-End Test Flows")

    content_slide(
        prs,
        "E2E Flow A — Full Setup to First Sale",
        [
            "1. Register → Admin approves → Login",
            "2. Create business profile (name, phone, GST)",
            "3. Settings: add tax, invoice setting (year/prefix/counter)",
            "4. Catalog: unit, category, brand, vendor, payment type",
            "5. Product: create with SKU, category, unit, MRP → generate barcode",
            "6. Customer: name, mobile, address, 6-digit pin code",
            "7. Stock In: vendor, invoice#, add product line, cost, save → verify stock",
            "8. Create Invoice: customer, products, finalize → verify stock reduced",
            "9. Print preview → verify invoice number, GST, terms, Mob. No.",
            "10. Record partial payment → verify pending bill + next due date",
        ],
        font_size=14,
    )

    content_slide(
        prs,
        "E2E Flow B — Draft, Cancel & Edge Cases",
        [
            "1. Create draft sale → verify no stock change, no invoice number",
            "2. Edit draft (add/remove lines) → finalize → stock deducts",
            "3. Create sale → cancel with reason → stock restored",
            "4. Stock-in → partial sell → try delete stock-in → must FAIL",
            "5. Sell product → try delete product → must FAIL",
            "6. Sell product → edit MRP only → must PASS",
            "7. Invoice series at end_counter → next sale rolls to new year series",
        ],
        font_size=14,
    )

    content_slide(
        prs,
        "E2E Flow C — Barcode & Stock-In Integration",
        [
            "1. Create product without barcode",
            "2. Open Stock In → add product line → generate barcode modal opens",
            "3. Generate → barcode fills readonly field",
            "4. Save stock-in → batch created with auto batch number",
            "5. Go to Settings → Barcode → Reprint tab → select SKU → print",
            "6. Try delete barcode → FAIL (product used in stock-in)",
        ],
        font_size=14,
    )

    section_slide(prs, "12. UI Pages Reference")

    table_slide(
        prs,
        "Frontend Routes (Manual UI Testing)",
        ["Page", "URL Path"],
        [
            ("Login", "/login/"),
            ("Register", "/register/"),
            ("Dashboard", "/dashboard/"),
            ("Products", "/dashboard/products/"),
            ("Units", "/dashboard/units/"),
            ("Categories", "/dashboard/categories/"),
            ("Brands", "/dashboard/brands/"),
            ("Vendors", "/dashboard/vendors/"),
            ("Inventory", "/dashboard/inventory/"),
            ("Stock In", "/dashboard/stock-in/"),
            ("Create Invoice (Sales)", "/dashboard/purchases/"),
            ("Customers", "/dashboard/customers/"),
            ("Business Profile", "/dashboard/business/"),
            ("Tax Settings", "/dashboard/settings/tax/"),
            ("Invoice Settings", "/dashboard/settings/invoice/"),
            ("Barcode Settings", "/dashboard/settings/barcode/"),
        ],
    )

    section_slide(prs, "13. Regression Checklist")

    content_slide(
        prs,
        "Pre-Release Regression Checklist",
        [
            "☐ Auth: login/logout, inactive user block, password validation",
            "☐ Business profile save with logo upload/clear",
            "☐ All catalog CRUD + soft delete + restore",
            "☐ Product create/edit/delete rules (with/without sales)",
            "☐ Customer pin_code 6-digit validation",
            "☐ Stock-in create/edit/delete with attachment limits",
            "☐ Batch creation and MRP auto-bump logic",
            "☐ Sale draft → finalize → payment → cancel cycle",
            "☐ FIFO stock deduction and restoration on cancel",
            "☐ Invoice numbering and series rollover",
            "☐ Barcode one-per-SKU and delete protection",
            "☐ Print preview: terms edit, Mob. No., company block",
            "☐ Pagination, search, ordering on all list pages",
            "☐ Staff role tab permissions (403 on restricted tabs)",
        ],
        font_size=13,
    )

    title_slide(
        prs,
        "End of Tester Guide",
        "DailyBasket Billing System\nRefer to API docs and codebase for latest changes",
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUTPUT))
    print(f"Created: {OUTPUT}")
    return OUTPUT


if __name__ == "__main__":
    build()
