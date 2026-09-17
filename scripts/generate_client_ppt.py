"""
Generate client-facing PowerPoint for DailyBasket Billing System.
Run: python scripts/generate_client_ppt.py
Output: docs/DailyBasket_Billing_System_Client_Guide.pptx
"""

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

OUTPUT = Path(__file__).resolve().parent.parent / "docs" / "DailyBasket_Billing_System_Client_Guide.pptx"

# Brand colors
GREEN = RGBColor(0x2E, 0x7D, 0x32)
DARK = RGBColor(0x1A, 0x1A, 0x2E)
GRAY = RGBColor(0x55, 0x55, 0x55)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_BG = RGBColor(0xF5, 0xF9, 0xF5)


def set_slide_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_title_slide(prs, title, subtitle=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    set_slide_bg(slide, GREEN)

    box = slide.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(8.4), Inches(1.5))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER

    if subtitle:
        box2 = slide.shapes.add_textbox(Inches(0.8), Inches(3.8), Inches(8.4), Inches(1.2))
        tf2 = box2.text_frame
        p2 = tf2.paragraphs[0]
        p2.text = subtitle
        p2.font.size = Pt(20)
        p2.font.color.rgb = WHITE
        p2.alignment = PP_ALIGN.CENTER


def add_section_slide(prs, section_title):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, DARK)
    box = slide.shapes.add_textbox(Inches(0.8), Inches(3.0), Inches(8.4), Inches(1.2))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.text = section_title
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER


def add_content_slide(prs, title, bullets, note=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, LIGHT_BG)

    # Title bar
    bar = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(10), Inches(1.1))
    bar.fill.solid()
    bar.fill.fore_color.rgb = GREEN
    bar.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.15), Inches(9), Inches(0.8))
    tp = title_box.text_frame.paragraphs[0]
    tp.text = title
    tp.font.size = Pt(28)
    tp.font.bold = True
    tp.font.color.rgb = WHITE

    body = slide.shapes.add_textbox(Inches(0.6), Inches(1.4), Inches(8.8), Inches(5.0))
    tf = body.text_frame
    tf.word_wrap = True

    for i, item in enumerate(bullets):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        if isinstance(item, tuple):
            text, level = item
            p.text = text
            p.level = level
        else:
            p.text = item
            p.level = 0
        p.font.size = Pt(18 if p.level == 0 else 16)
        p.font.color.rgb = DARK
        p.space_after = Pt(10)

    if note:
        note_box = slide.shapes.add_textbox(Inches(0.6), Inches(6.5), Inches(8.8), Inches(0.6))
        np = note_box.text_frame.paragraphs[0]
        np.text = note
        np.font.size = Pt(14)
        np.font.italic = True
        np.font.color.rgb = GRAY


def add_flow_slide(prs, title, steps):
    """Numbered step-by-step flow slide."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, LIGHT_BG)

    bar = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(10), Inches(1.1))
    bar.fill.solid()
    bar.fill.fore_color.rgb = GREEN
    bar.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.15), Inches(9), Inches(0.8))
    tp = title_box.text_frame.paragraphs[0]
    tp.text = title
    tp.font.size = Pt(26)
    tp.font.bold = True
    tp.font.color.rgb = WHITE

    y = 1.35
    for num, step_title, step_desc in steps:
        # Step number circle area (text)
        num_box = slide.shapes.add_textbox(Inches(0.5), Inches(y), Inches(0.5), Inches(0.4))
        np = num_box.text_frame.paragraphs[0]
        np.text = str(num)
        np.font.size = Pt(20)
        np.font.bold = True
        np.font.color.rgb = GREEN

        step_box = slide.shapes.add_textbox(Inches(1.1), Inches(y), Inches(8.3), Inches(0.9))
        stf = step_box.text_frame
        stf.word_wrap = True
        p1 = stf.paragraphs[0]
        p1.text = step_title
        p1.font.size = Pt(17)
        p1.font.bold = True
        p1.font.color.rgb = DARK
        p2 = stf.add_paragraph()
        p2.text = step_desc
        p2.font.size = Pt(14)
        p2.font.color.rgb = GRAY
        y += 0.85


def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    # --- Title ---
    add_title_slide(
        prs,
        "DailyBasket Billing System",
        "Complete Business Flow Guide\nFor Shop Owners & Staff",
    )

    # --- What is this system ---
    add_content_slide(
        prs,
        "What Is This System?",
        [
            "A complete shop management solution for your retail business",
            "Helps you manage products, stock, customers, suppliers, and billing — all in one place",
            "Designed for daily use: from receiving goods to printing customer invoices",
            "No technical knowledge needed — just follow the menus on screen",
        ],
        note="Example business: DailyBasket (your shop name appears on all invoices)",
    )

    add_section_slide(prs, "Part A — Getting Started")

    add_flow_slide(
        prs,
        "Step 1: Login & Dashboard",
        [
            (1, "Open the website", "Go to your billing system URL and sign in with your username and password."),
            (2, "Select your business", "If you manage more than one shop, choose the active business (e.g. DailyBasket)."),
            (3, "View the Dashboard", "See a quick summary — recent activity, stock alerts, and shortcuts to common tasks."),
            (4, "Use the side menu", "Navigate to Products, Customers, Stock In, Create Invoice, Settings, and more."),
        ],
    )

    add_flow_slide(
        prs,
        "Step 2: Initial Setup (Settings)",
        [
            (1, "Business Profile", "Add shop name, address, mobile number, email, and GST details — these appear on printed invoices."),
            (2, "Invoice Settings", "Set invoice prefix, terms & conditions, and default notes for customer bills."),
            (3, "Tax Settings", "Configure GST rates (CGST, SGST, IGST) as per your business requirements."),
            (4, "Barcode Setup", "Generate one barcode per product (SKU). Used for quick scanning during stock-in and billing."),
        ],
    )

    add_section_slide(prs, "Part B — Master Data Setup")

    add_flow_slide(
        prs,
        "Step 3: Add Product Categories",
        [
            (1, "Go to Categories", "Open the Categories section from the inventory menu."),
            (2, "Create a category", "Enter category name (e.g. Hair Care, Skin Care, Groceries)."),
            (3, "Save", "Categories help you organize products and filter reports later."),
        ],
    )

    add_flow_slide(
        prs,
        "Step 4: Add Products to Inventory",
        [
            (1, "Go to Products", "Open the Products page from the side menu."),
            (2, "Click Add Product", "Fill in product name, SKU number, category, unit, and MRP."),
            (3, "Set pricing", "Enter MRP (Maximum Retail Price) and optional salon/selling price."),
            (4, "Generate barcode", "Each product gets one unique barcode — used for scanning and labels."),
            (5, "Save product", "Product is now ready for stock-in and sales. Stock quantity starts at zero until you receive goods."),
        ],
    )

    add_flow_slide(
        prs,
        "Step 5: Add Customers",
        [
            (1, "Go to Customers", "Open the Customers section."),
            (2, "Click Add Customer", "Enter customer name, mobile number, and 6-digit PIN code (required)."),
            (3, "Company details (optional)", "If the customer is a business, add company name, company mobile, and email."),
            (4, "Address", "Add billing/delivery address for invoices."),
            (5, "Save", "Customer appears in the list — select them when creating a sale invoice."),
        ],
    )

    add_flow_slide(
        prs,
        "Step 6: Add Vendors (Suppliers)",
        [
            (1, "Go to Vendors", "Open the Vendors section."),
            (2, "Add supplier details", "Enter vendor name, contact number, and address."),
            (3, "Save", "Vendors are selected when you record stock received from suppliers (Stock In)."),
        ],
    )

    add_section_slide(prs, "Part C — Stock Management")

    add_content_slide(
        prs,
        "Understanding Stock & Batches",
        [
            "Stock = how many units of each product you currently have in the shop",
            "Every time you receive goods (Stock In), the system adds stock automatically",
            "Batch = a group of items received together at one cost price on one date",
            "Batch number is auto-generated (date + vendor + cost) — helps track different purchase prices",
            "When you sell, the system uses oldest stock first (FIFO) — fair and accurate costing",
        ],
    )

    add_flow_slide(
        prs,
        "Step 7: Stock In (Receive Goods from Supplier)",
        [
            (1, "Go to Stock In / Purchases", "Open the page where you record incoming inventory."),
            (2, "Select vendor", "Choose which supplier delivered the goods."),
            (3, "Add invoice details", "Enter supplier invoice number and date."),
            (4, "Add product lines", "For each item: select product, enter quantity, cost price with tax, and MRP."),
            (5, "Barcode auto-fills", "If the product has a barcode, it appears automatically. Generate one if missing."),
            (6, "Batch number", "Created automatically when you enter vendor and cost — or edit manually if needed."),
            (7, "Save Stock In", "Stock increases immediately. Each line creates a new batch for tracking."),
        ],
    )

    add_section_slide(prs, "Part D — Sales & Billing")

    add_flow_slide(
        prs,
        "Step 8: Create Sale Invoice (Sell to Customer)",
        [
            (1, "Go to Create Invoice / Sales", "Open the sales billing page."),
            (2, "Select customer", "Choose from your customer list or add a walk-in customer."),
            (3, "Add products", "Search by name or scan barcode. Enter quantity for each item."),
            (4, "Pricing", "System shows MRP and salon price. Adjust discount if needed."),
            (5, "Review totals", "See subtotal, GST breakdown (CGST/SGST/IGST), and grand total."),
            (6, "Payment", "Record amount paid — full or partial. Pending balance is tracked automatically."),
            (7, "Save invoice", "Stock decreases automatically (oldest batches used first). Invoice number is generated."),
        ],
    )

    add_flow_slide(
        prs,
        "Step 9: Print & Share Invoice",
        [
            (1, "Open saved invoice", "Find the invoice in the sales list and click to view."),
            (2, "Print Preview", "See exactly how the invoice will look on paper — shop details, customer info, item list, GST, total."),
            (3, "Edit Terms & Conditions", "Customize terms for this invoice before printing (optional)."),
            (4, "Print or Download", "Print for customer or save as PDF for WhatsApp/email."),
            (5, "Company customers", "Invoice shows company name, mobile (Mob. No.), and email when available."),
        ],
    )

    add_section_slide(prs, "Part E — Barcode & Labels")

    add_flow_slide(
        prs,
        "Step 10: Barcode Management",
        [
            (1, "One barcode per product", "Each SKU has exactly one barcode — no duplicates allowed."),
            (2, "Generate from Settings", "Go to Settings → Barcode → select product → Generate."),
            (3, "Generate during Stock In", "If a product has no barcode, the system prompts you to create one."),
            (4, "Reprint labels", "Go to Reprint tab → select SKU → print barcode label for shelf/packaging."),
            (5, "Delete protection", "Barcode cannot be deleted if the product was used on any purchase invoice."),
        ],
    )

    add_section_slide(prs, "Part F — Daily Workflow")

    add_content_slide(
        prs,
        "Typical Day — Start to End",
        [
            "Morning: Check Dashboard for low-stock alerts",
            "Receive delivery: Stock In → select vendor → add items → save (stock updated)",
            "Customer walk-in: Create Invoice → add products → take payment → print bill",
            "End of day: Review sales list, check pending payments from customers",
            "Weekly: Review product list, update MRP if supplier prices changed",
        ],
    )

    add_content_slide(
        prs,
        "Key Features at a Glance",
        [
            ("Products & Categories — Organize your entire catalog", 0),
            ("Stock In — Record purchases from suppliers, auto batch tracking", 0),
            ("Create Invoice — Bill customers with GST, discounts, partial payments", 0),
            ("Customers — Store contact, PIN code, company details", 0),
            ("Vendors — Track who supplies your goods", 0),
            ("Barcodes — Scan, print, and reprint product labels", 0),
            ("Print Preview — Professional invoices with editable terms", 0),
            ("FIFO Stock — Accurate inventory when costs vary by batch", 0),
        ],
    )

    add_content_slide(
        prs,
        "Important Terms (Simple Glossary)",
        [
            ("SKU — Your internal product code / item number", 0),
            ("MRP — Maximum Retail Price printed on the product", 0),
            ("Salon Price — Your actual selling price (may differ from MRP)", 0),
            ("Batch — One receipt of stock at a specific cost and date", 0),
            ("Stock In — Buying from supplier (stock goes UP)", 0),
            ("Sale / Invoice — Selling to customer (stock goes DOWN)", 0),
            ("GST — Tax shown separately on invoice (CGST + SGST or IGST)", 0),
            ("Pending Bill — Amount customer still owes you", 0),
        ],
    )

    # Closing
    add_title_slide(
        prs,
        "Thank You",
        "DailyBasket Billing System\nYour complete shop management partner",
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUTPUT))
    print(f"Created: {OUTPUT}")
    return OUTPUT


if __name__ == "__main__":
    build_presentation()
